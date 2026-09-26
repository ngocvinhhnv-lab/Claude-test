"""Tạo file mẫu (giả lập file xuất KiotViet) để chạy thử các công cụ.

  python3 -m tools.tao_du_lieu_mau
"""
import random
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

# mã, tên, SL bán TB/ngày, tồn kho, giá vốn, giá bán
SAN_PHAM = [
    ("LB-2027-01", "Lịch bloc 2027 cỡ đại 30x40", 6.0, 40, 38000, 89000),
    ("LB-2027-02", "Lịch bloc 2027 cỡ trung 20x30", 4.0, 300, 26000, 65000),
    ("LDB-2027-05", "Lịch để bàn 2027 chữ A", 3.0, 0, 21000, 55000),
    ("TR-THP-60", "Tranh thư pháp Phúc Lộc Thọ 60x90", 1.5, 25, 120000, 249000),
    ("LIEN-TET-01", "Liễn Tết gỗ treo tường", 0.0, 60, 55000, 135000),
    ("TRA-OL-500", "Trà ô long nguyên liệu 500g", 8.0, 500, 62000, 115000),
    ("TRA-XT-1KG", "Trà xanh lài nguyên liệu 1kg", 5.0, 90, 48000, 69000),
]
THU_MUC = Path("du_lieu_mau")


def tao(ngay_chot=datetime(2026, 9, 25), so_ngay=45, seed=1):
    rnd = random.Random(seed)
    THU_MUC.mkdir(exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.append(["BÁO CÁO CHI TIẾT HOÁ ĐƠN"])
    ws.append([f"Từ ngày ... đến ngày {ngay_chot:%d/%m/%Y}"])
    ws.append([])
    ws.append(["Mã hoá đơn", "Thời gian", "Mã hàng", "Tên hàng", "Số lượng", "Đơn giá", "Trạng thái"])
    stt = 0
    for i in range(so_ngay):
        ngay = ngay_chot - timedelta(days=so_ngay - 1 - i)
        for ma, ten, tb, _, _, gia in SAN_PHAM:
            # bán tăng dần về cuối kỳ để có xu hướng
            tb_ngay = tb * (0.7 + 0.6 * i / so_ngay)
            sl = max(0, round(rnd.gauss(tb_ngay, tb_ngay * 0.4)))
            if not sl:
                continue
            stt += 1
            gio = ngay.replace(hour=rnd.randint(8, 21), minute=rnd.randint(0, 59))
            trang_thai = "Đã huỷ" if rnd.random() < 0.05 else "Hoàn thành"
            ws.append([f"HD{stt:06d}", gio.strftime("%d/%m/%Y %H:%M:%S"), ma, ten, sl, gia, trang_thai])
    wb.save(THU_MUC / "ban_hang_mau.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.append(["Mã hàng", "Tên hàng", "Tồn kho"])
    for ma, ten, _, ton, _, _ in SAN_PHAM:
        ws.append([ma, ten, ton])
    wb.save(THU_MUC / "ton_kho_mau.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.append(["Mã hàng", "Tên hàng", "Giá vốn", "Giá bán"])
    for ma, ten, _, _, von, gia in SAN_PHAM:
        ws.append([ma, ten, von, gia])
    wb.save(THU_MUC / "san_pham_mau.xlsx")
    print(f"Đã tạo file mẫu trong thư mục {THU_MUC}/")


if __name__ == "__main__":
    tao()
