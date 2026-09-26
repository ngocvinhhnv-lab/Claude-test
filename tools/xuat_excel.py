"""Ghi bảng kết quả ra Excel có định dạng sẵn để in hoặc gửi nhân viên."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_TIEU_DE = PatternFill("solid", fgColor="1F4E78")
MAU = {
    "do": PatternFill("solid", fgColor="F8CBAD"),
    "cam": PatternFill("solid", fgColor="FFE699"),
    "xanh": PatternFill("solid", fgColor="C6EFCE"),
    "xam": PatternFill("solid", fgColor="D9D9D9"),
}


def ghi_sheet(ws, cot, dong, dinh_dang=None, mau_dong=None):
    """cot: [(tiêu đề, độ rộng)], dong: list các list giá trị.

    dinh_dang: {vị trí cột: number_format}; mau_dong: hàm(dòng) -> khoá trong MAU hoặc None.
    """
    dinh_dang = dinh_dang or {}
    ws.append([t for t, _ in cot])
    for i, (_, rong) in enumerate(cot, start=1):
        o = ws.cell(row=1, column=i)
        o.font = Font(bold=True, color="FFFFFF")
        o.fill = _TIEU_DE
        o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = rong
    ws.row_dimensions[1].height = 32
    for d in dong:
        ws.append(d)
        r = ws.max_row
        for vt, dd in dinh_dang.items():
            ws.cell(row=r, column=vt + 1).number_format = dd
        mau = mau_dong(d) if mau_dong else None
        if mau:
            for i in range(1, len(cot) + 1):
                ws.cell(row=r, column=i).fill = MAU[mau]
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def tao_file(duong_dan, cac_sheet):
    """cac_sheet: [(tên sheet, hàm(ws) ghi nội dung)]."""
    wb = Workbook()
    wb.remove(wb.active)
    for ten, ghi in cac_sheet:
        ghi(wb.create_sheet(ten))
    Path(duong_dan).parent.mkdir(parents=True, exist_ok=True)
    wb.save(duong_dan)
    return duong_dan
