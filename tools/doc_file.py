"""Đọc file Excel/CSV xuất từ KiotViet, Sapo, Shopee, TikTok Shop.

Tự tìm dòng tiêu đề (bỏ qua các dòng tên báo cáo phía trên) và nhận diện cột
theo nhiều cách đặt tên khác nhau của từng nền tảng.
"""
import csv
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

# Tên cột chuẩn -> các tên có thể gặp trong file xuất (so khớp không dấu, không phân biệt hoa thường)
BI_DANH_COT = {
    "ngay": [
        "thoi gian", "ngay", "ngay ban", "ngay dat hang", "thoi gian dat hang",
        "thoi gian tao", "ngay tao", "ngay hoa don", "ngay tao don",
        "order creation date", "created time", "order date",
    ],
    "ma": [
        "ma hang", "ma sku", "sku", "sku phan loai hang", "sku san pham",
        "ma san pham", "seller sku", "ma hang hoa", "sku nguoi ban", "ma",
    ],
    "ten": [
        "ten hang", "ten san pham", "ten hang hoa", "san pham", "product name",
        "ten phan loai hang", "ten",
    ],
    "so_luong": [
        "so luong", "sl", "sl ban", "so luong ban", "quantity", "sl ban ra",
    ],
    "ton": [
        "ton kho", "ton", "sl ton", "so luong ton", "ton cuoi ky", "stock",
        "co the ban", "ton hien tai",
    ],
    "trang_thai": [
        "trang thai", "trang thai don hang", "order status", "trang thai hoa don",
    ],
    "gia_von": ["gia von", "cost", "gia nhap"],
    "gia_ban": ["gia ban", "don gia", "price", "gia niem yet"],
}


def khong_dau(s):
    """'Mã Hàng' -> 'ma hang'."""
    s = str(s).replace("đ", "d").replace("Đ", "D")
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()


def nhan_dien_cot(tieu_de):
    """Trả về {tên chuẩn: vị trí cột} cho các cột nhận diện được."""
    chuan_hoa = [khong_dau(t) if t is not None else "" for t in tieu_de]
    ket_qua = {}
    for ten_chuan, bi_danh in BI_DANH_COT.items():
        for ten in bi_danh:  # ưu tiên theo thứ tự bí danh
            if ten in chuan_hoa:
                ket_qua[ten_chuan] = chuan_hoa.index(ten)
                break
    return ket_qua


def _doc_dong(duong_dan):
    duong_dan = Path(duong_dan)
    if duong_dan.suffix.lower() == ".csv":
        with open(duong_dan, encoding="utf-8-sig", newline="") as f:
            mau = f.read(4096)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(mau, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            return [list(d) for d in csv.reader(f, dialect)]
    wb = load_workbook(duong_dan, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    dong = [list(d) for d in ws.iter_rows(values_only=True)]
    wb.close()
    return dong


def doc_bang(duong_dan, cot_bat_buoc):
    """Đọc file, trả về danh sách dict theo tên cột chuẩn.

    Dò 20 dòng đầu để tìm dòng tiêu đề chứa đủ các cột bắt buộc.
    """
    dong = _doc_dong(duong_dan)
    for i, tieu_de in enumerate(dong[:20]):
        cot = nhan_dien_cot(tieu_de)
        if all(c in cot for c in cot_bat_buoc):
            break
    else:
        thieu = ", ".join(cot_bat_buoc)
        raise ValueError(
            f"{duong_dan}: không tìm thấy dòng tiêu đề có đủ cột [{thieu}]. "
            "Kiểm tra lại tên cột trong file (xem BI_DANH_COT trong tools/doc_file.py)."
        )
    ket_qua = []
    for d in dong[i + 1:]:
        if not any(v not in (None, "") for v in d):
            continue
        ket_qua.append({ten: d[vt] if vt < len(d) else None for ten, vt in cot.items()})
    return ket_qua


def doc_so(v):
    """Chuyển '1.200', '1,200', '₫120.000', 12.5 ... thành số. Rỗng -> 0."""
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.,\-]", "", str(v))
    if not s or s == "-":
        return 0.0
    if "." in s and "," in s:
        # dấu xuất hiện sau cùng là dấu thập phân
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif re.fullmatch(r"-?\d{1,3}([.,]\d{3})+", s):
        s = s.replace(".", "").replace(",", "")  # dấu phân cách hàng nghìn
    else:
        s = s.replace(",", ".")
    return float(s)


_DINH_DANG_NGAY = [
    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
    "%d-%m-%Y %H:%M:%S", "%d-%m-%Y",
]


def doc_ngay(v):
    """Chuyển ô ngày (datetime hoặc chuỗi dd/mm/yyyy, yyyy-mm-dd...) thành date."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if v is None:
        return None
    s = str(v).strip()
    for dd in _DINH_DANG_NGAY:
        try:
            return datetime.strptime(s, dd).date()
        except ValueError:
            pass
    return None
